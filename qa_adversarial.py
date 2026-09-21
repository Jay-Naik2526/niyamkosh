"""Try to break the API the way a judge or a careless officer would.

Nothing here checks that the system is clever. It checks that when the input is
wrong, the failure is a clear status code and a sentence a human can act on —
never a stack trace, never a 500, and never a confident answer to a question
that was never asked.

Run against a live server:  python qa_adversarial.py
"""

import io
import json
import sys
import urllib.error
import urllib.request
import uuid

BASE = "http://localhost:8000"

PASS, FAIL = [], []


def call(method: str, path: str, payload=None, raw: bytes | None = None,
         content_type: str = "application/json"):
    url = BASE + path
    data = raw if raw is not None else (json.dumps(payload).encode() if payload is not None else None)
    headers = {"Content-Type": content_type} if data else {}
    req = urllib.request.Request(url, data=data, headers=headers, method=method)
    try:
        with urllib.request.urlopen(req, timeout=120) as r:
            return r.status, json.loads(r.read() or b"null")
    except urllib.error.HTTPError as e:
        body = e.read()
        try:
            return e.code, json.loads(body or b"null")
        except json.JSONDecodeError:
            return e.code, {"raw": body[:200].decode("utf-8", "replace")}
    except Exception as exc:                       # noqa: BLE001 — report, don't crash the run
        return None, {"transport_error": f"{type(exc).__name__}: {exc}"}


def check(name: str, ok: bool, note: str = ""):
    (PASS if ok else FAIL).append(name)
    print(f"  {'PASS' if ok else 'FAIL'}  {name}" + (f"  — {note}" if note else ""))


def multipart(filename: str, content: bytes) -> tuple[bytes, str]:
    boundary = uuid.uuid4().hex
    body = (
        f"--{boundary}\r\nContent-Disposition: form-data; name=\"file\"; "
        f"filename=\"{filename}\"\r\nContent-Type: application/octet-stream\r\n\r\n"
    ).encode() + content + f"\r\n--{boundary}--\r\n".encode()
    return body, f"multipart/form-data; boundary={boundary}"


def section(title: str):
    print(f"\n{title}\n" + "-" * len(title))



def _fetch_report(cited: list[str]) -> str:
    """The report endpoint returns HTML, which `call` cannot parse as JSON."""
    import urllib.request

    data = json.dumps({"cited": cited, "document": "adversarial"}).encode()
    req = urllib.request.Request(BASE + "/report", data=data,
                                 headers={"Content-Type": "application/json"},
                                 method="POST")
    with urllib.request.urlopen(req, timeout=120) as r:
        return r.read().decode("utf-8", "replace")


def _plain(html_text: str) -> str:
    """Visible text only — the CSS mentions colours, not standards."""
    import html as _html
    import re as _re

    body = html_text.split("</style>", 1)[-1]
    return _re.sub(r"\s+", " ", _html.unescape(_re.sub(r"<[^>]+>", " ", body)))


def main():
    status, _ = call("GET", "/health")
    if status != 200:
        print(f"Server not answering at {BASE}. Start it first.")
        sys.exit(2)

    # ---------------------------------------------------------------- empties
    section("Empty and near-empty input")

    s, b = call("POST", "/analyze", {"cited_is_numbers": []})
    check("/analyze with nothing at all returns 200 and an empty result, not a guess",
          s == 200 and b.get("matched_standards") is None, f"status {s}")

    s, b = call("POST", "/audit-text", {"text": ""})
    check("/audit-text with empty text is rejected with a usable message",
          s == 400 and "IS number" in str(b.get("detail", "")), f"status {s}")

    s, b = call("POST", "/audit-text", {"text": "   \n\t  "})
    check("/audit-text with whitespace only is rejected", s == 400, f"status {s}")

    s, b = call("POST", "/recommend", {"spec_text": "   "})
    check("/recommend with blank spec is rejected, not embedded", s == 400, f"status {s}")

    s, b = call("POST", "/audit-text", {"text": "Supply of cables. No standards named."})
    check("/audit-text with prose but no citations says so instead of inventing one",
          s == 200 and b["cited_count"] == 0 and not b["findings"], f"status {s}")

    # ---------------------------------------------------------------- garbage
    section("Malformed and hostile citations")

    s, b = call("POST", "/analyze", {"cited_is_numbers": ["IS 99999999", "IS ABC", "", "IS -1"]})
    check("/analyze survives nonsense IS numbers and reports them as not found",
          s == 200 and all(not v.get("found") for v in b["dead_citations"].values()), f"status {s}")

    s, b = call("POST", "/audit-text", {"cited": ["'; DROP TABLE standards; --"]})
    check("SQL in a citation is parameterised, not executed", s == 200, f"status {s}")
    s2, _ = call("GET", "/health")
    check("standards table still exists after that", s2 == 200)

    s, b = call("POST", "/audit-text", {"text": "<script>alert(1)</script> IS 1554"})
    check("script tags pass through as inert text, escaped client-side",
          s == 200 and "IS 1554" in b["cited"], f"status {s}")

    s, b = call("GET", "/evidence?is_number=" + "%2E%2E%2F" * 6)
    check("path traversal in a query param returns an empty result, not a file",
          s == 200 and b["tenders_citing"] == 0, f"status {s}")

    # ---------------------------------------------------------------- scripts
    section("Non-Latin and mixed script")

    devanagari = (
        "भारतीय मानक "
        "IS 1554 के अनुसार "
        "केबल की आपूर्ति"
    )
    s, b = call("POST", "/audit-text", {"text": devanagari})
    check("Devanagari text with an embedded IS number extracts that citation",
          s == 200 and any("1554" in c for c in b["cited"]), f"status {s} cited {b.get('cited')}")

    s, b = call("POST", "/recommend", {"spec_text": devanagari})
    check("/recommend handles Devanagari without a 500", s == 200, f"status {s}")

    s, b = call("POST", "/audit-text", {"text": "\U0001f4a5\U0001f6d1 IS 694 — éèü"})
    check("emoji and accented Latin do not break extraction",
          s == 200 and any("694" in c for c in b["cited"]), f"status {s}")

    # ---------------------------------------------------------------- volume
    section("Volume")

    # Just under the stated limit — a genuinely long tender, not a probe.
    big = ("Supply of PVC insulated cables conforming to IS 1554. " * 7000)   # ~380 KB
    s, b = call("POST", "/audit-text", {"text": big})
    check("380 KB of repeated text is handled and de-duplicated to one citation",
          s == 200 and b["cited_count"] == 1, f"status {s} chars {b.get('characters')}")

    many = "; ".join(f"IS {n}" for n in range(1000, 1400))
    s, b = call("POST", "/audit-text", {"cited": many.split("; ")})
    check("400 distinct citations in one document are all resolved",
          s == 200 and len(b["resolved"]) == 400, f"status {s}")

    s, b = call("POST", "/recommend", {"spec_text": "cable " * 20000})
    check("/recommend clips a 120 KB query and says how much it actually read",
          s == 200 and b["input"]["truncated"] and b["input"]["characters_read"] < 3000,
          f"status {s} read {b.get('input', {}).get('characters_read')}")

    s, b = call("POST", "/audit-text", {"text": "x" * 500_000})
    check("half a megabyte of pasted text is refused with a stated limit",
          s == 413 and "characters" in str(b.get("detail", "")), f"status {s}")

    # ---------------------------------------------------------------- uploads
    section("Uploads")

    body, ct = multipart("empty.pdf", b"")
    s, b = call("POST", "/extract", raw=body, content_type=ct)
    check("empty upload is rejected with 400", s == 400, f"status {s}")

    body, ct = multipart("notes.txt", b"this is plain text, not a document")
    s, b = call("POST", "/extract", raw=body, content_type=ct)
    check("a .txt renamed as an upload is refused with an instruction",
          s == 400 and "PDF" in str(b.get("detail", "")), f"status {s}")

    body, ct = multipart("fake.pdf", b"%PDF-1.4 truncated garbage not a real pdf")
    s, b = call("POST", "/extract", raw=body, content_type=ct)
    check("a corrupt PDF returns 422 with a reason, not a traceback",
          s == 422 and "detail" in b, f"status {s}")

    scanned = _blank_pdf()
    if scanned:
        body, ct = multipart("scan.pdf", scanned)
        s, b = call("POST", "/audit-tender", raw=body, content_type=ct)
        check("a text-free PDF is reported as unreadable, never as clean",
              s == 200 and b["verdict"] == "unreadable" and b["extraction"]["scanned"],
              f"status {s} verdict {b.get('verdict')}")
    else:
        print("  SKIP  scanned-PDF case (reportlab/pypdf not installed)")

    # ---------------------------------------------------------------- decisions
    section("Decision log")

    s, b = call("POST", "/decision", {"officer": "QA", "decision": "overridden", "is_number": "IS 1"})
    check("an override with no rationale is refused",
          s == 400 and "rationale" in str(b.get("detail", "")).lower(), f"status {s}")

    s, b = call("POST", "/decision", {"officer": "QA", "decision": "banana"})
    check("an unknown decision verb is refused", s == 400, f"status {s}")

    section("Phantom citations")

    # A number one digit from a real standard, in a document whose other
    # citations the graph ties to that standard. 81304 is absent from the
    # catalogue; IS 6994 was used here until the harvest made it a real standard.
    s, b = call("POST", "/audit-text", {"text":
        "Cables shall conform to IS 81304. Conductors as per IS 8130 and insulation IS 5831."})
    hit = next((f for f in (b.get("findings") or [])
                if f.get("kind") == "not_in_register"), None)
    guess = (hit or {}).get("did_you_mean")
    check("a one-digit slip is offered as a question, with graph evidence",
          s == 200 and guess is not None
          and str(guess.get("is_number", "")).startswith("IS 8130")
          and bool(guess.get("shares_citations_with")),
          f"status {s} · suggested {(guess or {}).get('is_number')}")

    # No neighbours, no suggestion: distance alone must never be enough.
    s, b = call("POST", "/audit-text", {"text": "Material shall conform to IS 99999."})
    hit = next((f for f in (b.get("findings") or [])
                if f.get("kind") == "not_in_register"), None)
    check("a number with no graph support gets no suggestion",
          s == 200 and hit is not None and hit.get("did_you_mean") is None,
          f"status {s}")

    section("Peer citations")

    s, b = call("GET", "/peers?text=11%20kV%20XLPE%20cable%20joint")
    cites = (b or {}).get("citations") or []
    check("a real item finds a peer group and tallies its citations",
          s == 200 and b.get("found") and b.get("matched_documents", 0) >= 3
          and all(c["documents"] <= c["of"] for c in cites),
          f"status {s} · {b.get('matched_documents')} bids")

    # One weak word-overlap is not a peer group; saying nothing beats a corpus
    # average dressed up as a recommendation.
    s, b = call("GET", "/peers?text=banana")
    check("an item with no comparable bids returns nothing, not the corpus average",
          s == 200 and not b.get("found"), f"status {s} · found={b.get('found')}")

    s, b = call("GET", "/peers?text=")
    check("/peers with empty text is rejected", s == 400, f"status {s}")

    section("Compliance report")

    text = _plain(_fetch_report(["IS 434 (Part 1)", "IS 1570", "IS 694", "IS 999999"]))
    check("a report renders, names the supersession and its successor",
          "IS 434 (Part 1)" in text and "Superseded" in text and "IS 9968" in text
          and "not a certification of the tender" in text,
          f"{len(text)} characters of visible text")

    # The report reads the register's status per citation. It once reported
    # IS 694 — current, with a live QCO — as "not in the register", because a
    # certification finding carries no status field and an absent field was
    # read as absence from the register.
    text = _plain(_fetch_report(["IS 694"]))
    check("a current standard with a certification duty is not called missing",
          "IS 694" in text and "not in the register" not in text,
          "IS 694 " + ("mislabelled" if "not in the register" in text else "reported as held"))

    text = _plain(_fetch_report(["IS 999999"]))
    check("a number the register does not hold is reported as unresolved",
          "not in the register" in text, "")

    s, b = call("POST", "/report", {"text": "", "cited": []})
    check("a report with neither text nor citations is refused", s == 400, f"status {s}")

    # ---------------------------------------------------------------- summary
    # ------------------------------------------------- the repaired document
    section("Document repair: what comes back out")
    _check_repair()

    print(f"\n{len(PASS)} passed · {len(FAIL)} failed")
    if FAIL:
        print("\nfailures:")
        for f in FAIL:
            print(f"  - {f}")
        sys.exit(1)


def _download(token: str, fmt: str) -> bytes:
    import urllib.request
    with urllib.request.urlopen(BASE + f"/fix-download/{token}?fmt={fmt}", timeout=120) as r:
        return r.read()


def _check_repair():
    """The repaired file is the deliverable, so it is checked as a file.

    Both faults these assert against were shipping. The first pass rewrites the
    page's own text layer and the second pass overlays whatever the first could
    not reach — but a withdrawal note still begins with the citation it marks,
    so the second pass matched its own first pass and drew the note on top of
    itself, across the next column. And crossing between .docx and .pdf used to
    drop the letterhead entirely, which makes a tender that cannot be submitted.
    """
    import os

    for source in ("pdf", "docx"):
        path = os.path.join("samples", f"sample-tender.{source}")
        if not os.path.exists(path):
            check(f"sample-tender.{source} is present to test against", False, "missing")
            continue
        with open(path, "rb") as fh:
            body, ctype = multipart(os.path.basename(path), fh.read())
        s, b = call("POST", "/fix-document", raw=body, content_type=ctype)
        if s != 200 or not isinstance(b, dict) or not b.get("token"):
            check(f"/fix-document accepts a .{source} tender", False, f"status {s}")
            continue
        check(f"/fix-document accepts a .{source} tender", True)

        same = _download(b["token"], source)
        other = "docx" if source == "pdf" else "pdf"
        cross = _download(b["token"], other)

        if source == "pdf":
            import pdf_fix
            text = pdf_fix.pdf_text(same)
            # The invariant, not the symptom: a repair rewrites citations, it
            # never multiplies them. When the overlay pass re-patched text the
            # stream pass had already corrected, the note was drawn a second
            # time at its own baseline and the page came back carrying two of
            # every marked citation. Counted on page one of each file, because
            # the addendum page names the marked standards once by design.
            def _page_one(data: bytes) -> str:
                for i, _w, _h, words in pdf_fix.pdf_words(data, max_pages=1):
                    if i == 0:
                        return "".join(w["text"] for w in words).replace(" ", "")
                return ""

            with open(path, "rb") as fh:
                before = _page_one(fh.read())
            after = _page_one(same)
            flagged = [c["from"] for c in b.get("changes", []) if c.get("action") == "flag"]
            worst = ""
            for is_number in flagged:
                key = is_number.replace(" ", "")
                if after.count(key) > before.count(key):
                    worst = f"{is_number}: {before.count(key)} before, {after.count(key)} after"
            check("the repair rewrites citations without multiplying them",
                  bool(flagged) and not worst,
                  worst or ("no flag change in this sample" if not flagged else ""))
            import docx as _docx
            d = _docx.Document(io.BytesIO(cross))
            head = " ".join(p.text for p in d.sections[0].header.paragraphs).strip()
            check("pdf -> docx carries the letterhead into the Word header",
                  bool(head), f"header is {head[:40]!r}")
            body_text = "\n".join(p.text for p in d.paragraphs)
            check("the letterhead is not also printed at the top of the body",
                  not body_text.lstrip().startswith(head.split()[0] if head else "\0"))
        else:
            import pdf_fix
            text = pdf_fix.pdf_text(cross)
            check("docx -> pdf carries the letterhead onto the page",
                  "SUPERINTENDING ENGINEER" in text.upper(), text[:60])
            check("docx -> pdf keeps the schedule's rows",
                  text.count("IS 1554") >= 1 and "Governing standard" in text)

        notes = b.get("downloads", {}).get(other, {})
        check(f"the {other} download says whether the letterhead came with it",
              "letterhead_carried" in notes)


def _blank_pdf() -> bytes | None:
    """A one-page PDF with no text layer — what a scanned tender looks like to
    a parser. Built here rather than committed as a binary fixture."""
    try:
        import pypdf
    except ImportError:
        return None
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=595, height=842)
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


if __name__ == "__main__":
    main()
